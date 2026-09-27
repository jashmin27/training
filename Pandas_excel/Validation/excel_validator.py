"""
excel_validator.py

A small reusable utility to check an Excel sheet for common data problems:
- required columns missing values
- values outside an allowed range
- duplicate keys in a column that's supposed to be unique

How to use it:
    Edit the CONFIG section below to match whatever sheet you're checking
    (which columns are required, which columns have min/max limits, which
    column should have unique values), then run:

        python excel_validator.py employee_data.xlsx

    It prints every problem it finds, with the row number and what's wrong,
    and also writes them out to validation_report.txt so you have a copy.

You can also import validate_sheet() and call it directly if you're checking
files from another script instead of the command line.
"""

import sys
from openpyxl import load_workbook


# ---------------- CONFIG: edit this part for your own data ----------------

SHEET_NAME = "Employees"

REQUIRED_COLUMNS = ["EmployeeID", "Name", "Age", "Salary"]

RANGES = {
    "Age": (18, 65),
    "Salary": (0, None),   # no upper limit, just can't be negative
}

UNIQUE_COLUMN = "EmployeeID"

# ---------------------------------------------------------------------------


def validate_sheet(filepath, sheet_name=SHEET_NAME, required_columns=None,
                    ranges=None, unique_column=None):
    """
    Opens the workbook and checks every row against the rules above.
    Returns a list of issue strings. An empty list means the file is clean.
    """
    if required_columns is None:
        required_columns = REQUIRED_COLUMNS
    if ranges is None:
        ranges = RANGES
    if unique_column is None:
        unique_column = UNIQUE_COLUMN

    issues = []

    wb = load_workbook(filepath, data_only=True)

    if sheet_name not in wb.sheetnames:
        issues.append(f"Sheet '{sheet_name}' was not found in this file. "
                       f"Sheets available: {wb.sheetnames}")
        return issues

    ws = wb[sheet_name]

    header_row = [cell.value for cell in ws[1]]

    # Check all required columns actually exist in the header
    missing_headers = [col for col in required_columns if col not in header_row]
    if missing_headers:
        issues.append(f"Missing required column(s) in header row: {missing_headers}")
        return issues  # no point checking rows if the columns aren't even there

    col_index = {name: header_row.index(name) + 1 for name in header_row if name}

    seen_keys = {}

    for row_num in range(2, ws.max_row + 1):
        row_values = [cell.value for cell in ws[row_num]]

        # skip completely empty rows
        if all(v is None or str(v).strip() == "" for v in row_values):
            continue

        # required columns must not be blank
        for col in required_columns:
            cell_value = ws.cell(row=row_num, column=col_index[col]).value
            if cell_value is None or str(cell_value).strip() == "":
                issues.append(f"Row {row_num}: '{col}' is required but is blank")

        # range checks
        for col, (min_val, max_val) in ranges.items():
            if col not in col_index:
                continue
            cell_value = ws.cell(row=row_num, column=col_index[col]).value
            if cell_value is None:
                continue  # already flagged above if it's required
            try:
                num = float(cell_value)
            except (TypeError, ValueError):
                issues.append(f"Row {row_num}: '{col}' value '{cell_value}' is not a number")
                continue
            if min_val is not None and num < min_val:
                issues.append(f"Row {row_num}: '{col}' = {cell_value} is below the allowed "
                               f"minimum of {min_val}")
            if max_val is not None and num > max_val:
                issues.append(f"Row {row_num}: '{col}' = {cell_value} is above the allowed "
                               f"maximum of {max_val}")

        # duplicate key check
        if unique_column and unique_column in col_index:
            key_value = ws.cell(row=row_num, column=col_index[unique_column]).value
            if key_value is not None:
                if key_value in seen_keys:
                    issues.append(f"Row {row_num}: '{unique_column}' = '{key_value}' is a "
                                   f"duplicate (first seen in row {seen_keys[key_value]})")
                else:
                    seen_keys[key_value] = row_num

    return issues


def main():
    if len(sys.argv) < 2:
        print("Usage: python excel_validator.py <file.xlsx>")
        sys.exit(1)

    filepath = sys.argv[1]
    issues = validate_sheet(filepath)

    report_lines = [f"Validation report for: {filepath}", ""]

    if not issues:
        report_lines.append("No problems found. File looks good.")
    else:
        report_lines.append(f"{len(issues)} problem(s) found:")
        report_lines.append("")
        for issue in issues:
            report_lines.append(f"- {issue}")

    report_text = "\n".join(report_lines)
    print(report_text)

    with open("validation_report.txt", "w") as f:
        f.write(report_text + "\n")


if __name__ == "__main__":
    main()
