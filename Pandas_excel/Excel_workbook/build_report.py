"""
Small script to build a simple multi-sheet sales report.
Sheets: Sales, Expenses, Summary
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

wb = Workbook()

HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True)
NORMAL_FONT = Font(name="Arial")
TITLE_FONT = Font(name="Arial", bold=True, size=14)


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------- Sales sheet ----------------
ws = wb.active
ws.title = "Sales"

sales_data = [
    ("Date", "Region", "Product", "Units Sold", "Unit Price", "Revenue"),
    ("2026-01-05", "North", "Widget A", 40, 12.5, None),
    ("2026-01-12", "South", "Widget B", 25, 18.0, None),
    ("2026-01-18", "North", "Widget A", 30, 12.5, None),
    ("2026-01-22", "East", "Widget C", 15, 22.0, None),
    ("2026-02-03", "South", "Widget A", 50, 12.5, None),
    ("2026-02-10", "West", "Widget B", 35, 18.0, None),
    ("2026-02-14", "East", "Widget C", 20, 22.0, None),
    ("2026-02-25", "North", "Widget B", 28, 18.0, None),
    ("2026-03-02", "West", "Widget A", 45, 12.5, None),
    ("2026-03-09", "South", "Widget C", 18, 22.0, None),
]

for r, row in enumerate(sales_data, start=1):
    for c, val in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=val)

# Revenue = Units Sold * Unit Price, as a real formula
last_row = len(sales_data)
for r in range(2, last_row + 1):
    ws.cell(row=r, column=6, value=f"=D{r}*E{r}")

style_header(ws, 1, 6)
autofit(ws, [12, 10, 12, 12, 11, 12])
for r in range(2, last_row + 1):
    ws.cell(row=r, column=5).number_format = '$#,##0.00'
    ws.cell(row=r, column=6).number_format = '$#,##0.00'
ws.freeze_panes = "A2"

# ---------------- Expenses sheet ----------------
ws2 = wb.create_sheet("Expenses")

expense_data = [
    ("Date", "Category", "Description", "Amount"),
    ("2026-01-08", "Shipping", "Courier - North region", 120.00),
    ("2026-01-15", "Supplies", "Packaging boxes", 45.50),
    ("2026-01-29", "Marketing", "Local flyers", 80.00),
    ("2026-02-05", "Shipping", "Courier - South region", 150.00),
    ("2026-02-11", "Supplies", "Packaging tape", 20.00),
    ("2026-02-20", "Marketing", "Social media ads", 100.00),
    ("2026-03-04", "Shipping", "Courier - West region", 95.00),
    ("2026-03-10", "Supplies", "Labels", 30.00),
]

for r, row in enumerate(expense_data, start=1):
    for c, val in enumerate(row, start=1):
        ws2.cell(row=r, column=c, value=val)

exp_last_row = len(expense_data)
style_header(ws2, 1, 4)
autofit(ws2, [12, 12, 24, 12])
for r in range(2, exp_last_row + 1):
    ws2.cell(row=r, column=4).number_format = '$#,##0.00'
ws2.freeze_panes = "A2"

# ---------------- Summary sheet ----------------
ws3 = wb.create_sheet("Summary")
ws3["A1"] = "Sales Report Summary"
ws3["A1"].font = TITLE_FONT
ws3["A2"] = "Q1 2026 (Jan - Mar)"
ws3["A2"].font = Font(name="Arial", italic=True)

# Totals block
ws3["A4"] = "Metric"
ws3["B4"] = "Value"
style_header(ws3, 4, 2)

ws3["A5"] = "Total Revenue"
ws3["B5"] = f"=SUM(Sales!F2:F{last_row})"
ws3["B5"].number_format = '$#,##0.00'

ws3["A6"] = "Total Expenses"
ws3["B6"] = f"=SUM(Expenses!D2:D{exp_last_row})"
ws3["B6"].number_format = '$#,##0.00'

ws3["A7"] = "Net Profit"
ws3["B7"] = "=B5-B6"
ws3["B7"].number_format = '$#,##0.00'
ws3["A7"].font = Font(name="Arial", bold=True)
ws3["B7"].font = Font(name="Arial", bold=True)

ws3["A8"] = "Total Units Sold"
ws3["B8"] = f"=SUM(Sales!D2:D{last_row})"

ws3["A9"] = "Average Revenue per Sale"
ws3["B9"] = f"=AVERAGE(Sales!F2:F{last_row})"
ws3["B9"].number_format = '$#,##0.00'

for r in range(5, 10):
    ws3.cell(row=r, column=1).font = NORMAL_FONT

# Revenue by region (using SUMIFS so it's a real formula, not hardcoded)
ws3["A11"] = "Revenue by Region"
ws3["A11"].font = Font(name="Arial", bold=True)

regions = ["North", "South", "East", "West"]
ws3["A12"] = "Region"
ws3["B12"] = "Revenue"
style_header(ws3, 12, 2)

for i, region in enumerate(regions, start=13):
    ws3.cell(row=i, column=1, value=region)
    ws3.cell(row=i, column=2,
              value=f'=SUMIFS(Sales!F2:F{last_row},Sales!B2:B{last_row},A{i})')
    ws3.cell(row=i, column=2).number_format = '$#,##0.00'

# Expenses by category
ws3["D11"] = "Expenses by Category"
ws3["D11"].font = Font(name="Arial", bold=True)

categories = ["Shipping", "Supplies", "Marketing"]
ws3["D12"] = "Category"
ws3["E12"] = "Amount"
style_header(ws3, 12, 5)  # harmless re-style, covers columns A-E already partly done above

ws3["D12"] = "Category"
ws3["E12"] = "Amount"
ws3["D12"].font = HEADER_FONT
ws3["D12"].fill = HEADER_FILL
ws3["E12"].font = HEADER_FONT
ws3["E12"].fill = HEADER_FILL

for i, cat in enumerate(categories, start=13):
    ws3.cell(row=i, column=4, value=cat)
    ws3.cell(row=i, column=5,
              value=f'=SUMIFS(Expenses!D2:D{exp_last_row},Expenses!B2:B{exp_last_row},D{i})')
    ws3.cell(row=i, column=5).number_format = '$#,##0.00'

autofit(ws3, [20, 14, 4, 18, 14])

wb.save("sales_report.xlsx")
print("saved")
